"""Export a static Presenton companion from final rendered pages; never upload source PDFs."""
from pathlib import Path
import tempfile,json,requests,concurrent.futures,html,hashlib
import fitz
from pptx import Presentation
D=Path(__file__).resolve().parent;PDF=D/'Sensing_the_Body_and_the_Chemical_World_Chapter_7_Static.pdf';PPTX=D/'Sensing_the_Body_and_the_Chemical_World_Chapter_7.pptx';API='https://api.presenton.ai';TITLE='Sensing the Body and the Chemical World: Balance, Touch, Pain, Taste, and Smell'
def response(r):
 r.raise_for_status();a=r.json()
 if isinstance(a,dict) and a.get('message'):print('API message: '+a['message'],flush=True)
 return a
with tempfile.TemporaryDirectory(prefix='presenton-') as directory:
 tmp=Path(directory);doc=fitz.open(PDF);prs=Presentation(PPTX);assert len(doc)==len(prs.slides)==75
 files=[]
 for i,p in enumerate(doc):
  f=tmp/f'slide-{i+1:02}.png';p.get_pixmap(matrix=fitz.Matrix(1280/p.rect.width,720/p.rect.height)).save(f);files.append(f)
 def upload(f):
  with f.open('rb') as stream:a=response(requests.post(API+'/api/v3/images/upload/public',files={'file':(f.name,stream,'image/png')},timeout=90))
  assert a['url'].startswith('https://');return a['url']
 urls=list(concurrent.futures.ThreadPoolExecutor(4).map(upload,files));print('Uploaded',len(urls),'final static slide images.',flush=True)
 slides=[]
 for i,u in enumerate(urls):
  notes=prs.slides[i].notes_slide.notes_text_frame.text
  slides.append(f'<section class="relative h-[720px] w-[1280px] overflow-hidden font-[Arial]" data-speaker-note="{html.escape(notes,quote=True)}"><img class="absolute left-0 top-0 h-[720px] w-[1280px] object-contain" src="{html.escape(u,quote=True)}" alt="Chapter 7 static slide {i+1}; original figures and source credits" width="1280" height="720"></section>')
 markup='<!doctype html><html lang="en"><head><meta charset="utf-8"><script src="https://cdn.tailwindcss.com"></script></head><body class="m-0 p-0"><main id="presentation-slides-wrapper" class="m-0 w-[1280px] p-0">'+''.join(slides)+'</main></body></html>'
 assert markup.count('<section ')==75 and 'data:' not in markup and '<style' not in markup and ' style=' not in markup
 (tmp/'presentation.html').write_text(markup)
 a=response(requests.post(API+'/api/v3/export/html-to-any',json={'html':markup,'format':'pdf','title':TITLE+' — Static Companion'},timeout=300));assert isinstance(a['id'],int) and a['id']>0;assert a['url'].startswith('https://');print('EXPORT',json.dumps(a),flush=True)
 preview=response(requests.post(API+'/api/v3/export/html-to-any/create-preview',json={'id':a['id']},timeout=90));assert preview['url'].startswith('https://');print('PREVIEW',preview['url'],flush=True)
 result={'title':TITLE,'slides':75,'format':'pdf','export':a,'preview':preview,'fonts':['Arial'],'native_pdf_sha256':hashlib.sha256(PDF.read_bytes()).hexdigest(),'source':'Final native static PDF pages, with original credits; full source book and article PDFs not uploaded.','expires':'Presenton export and preview links expire after 24 hours.'}
 (D/'presenton_preview.json').write_text(json.dumps(result,indent=2))
 print('Fonts in final HTML: Arial. Inventory applies to this static PDF companion. Temporary HTML and rendered pages cleaned automatically.',flush=True)
