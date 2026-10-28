"""Presenton static companions. Offline native movies remain in course-built PPTXs."""
from pathlib import Path
import sys,json,html,tempfile,re,concurrent.futures
from pptx import Presentation
sys.path.insert(0,'/workspace/.setup/presenton')
from presenton_artifacts import request_public_image_upload,request_json,print_response_message,print_font_usage
from validate_html import validate_html
R=Path(__file__).resolve().parent.parent
for ch in map(int,sys.argv[1:] or [18,19]):
 D=R/'chapters/CTE_Psychology';sp=json.load(open(D/'lecture.json'));cfg=json.load(open(D/'settings.json'));figs=json.load(open(D/'figure_sources.json'));media=json.load(open(D/'media_sources.json'))['sources'];by={m['slide']:m for m in media};assets=json.load(open(D/'presenton_assets.json')) if (D/'presenton_assets.json').exists() else {};urls={f['path']:f['original_url'] for f in figs.values()};prs=Presentation(D/sp['output_filename'])
 for m in media:
  if m['paper_id'].isdigit():
   j=json.load(open('/workspace/.setup/publisher/'+m['paper_id']+'.json'))
   def walk(n):
    if isinstance(n,dict):
     yield n
     for v in n.values():yield from walk(v)
    elif isinstance(n,list):
     for v in n:yield from walk(v)
   n=next(n for n in walk(j['body']) if n.get('type')=='video' and n.get('id')==m['identifier']);urls[m['poster']]=n['image']
 paths=[sp['title_image']['path']]+[s['figure']['path'] for s in sp['slides']]+[m['poster'] for m in media];needed=sorted(set(paths)-set(urls)-set(assets))
 def upload(path):
  a=request_public_image_upload(D/path,timeout=180);return path,a
 with concurrent.futures.ThreadPoolExecutor(4) as ex:
  for path,a in ex.map(upload,needed):print_response_message(a);assert a['url'].startswith('https://');assets[path]=a['url'];(D/'presenton_assets.json').write_text(json.dumps(assets,indent=2)+'\n')
 def url(p):return urls.get(p,assets.get(p))
 def e(t):return html.escape(str(t),quote=True)
 ink='#'+cfg.get('color','3D4A55');slides=[];title=sp['title_image']
 def section(body,note):return '<section class="relative h-[720px] w-[1280px] overflow-hidden bg-white font-[Arial]" data-speaker-note="'+e(note)+'">'+body+'</section>'
 slides.append(section(f'<img src="{e(url(title["path"]))}" alt="{e(title["caption"])}" width="530" height="465" class="absolute left-[60px] top-[115px] h-[465px] w-[530px] object-contain"><div class="absolute left-[650px] top-[85px] w-[560px] text-[17px] text-black">BEHAVIORAL NEUROSCIENCE 392</div><h1 class="absolute left-[650px] top-[150px] w-[560px] text-[39px] leading-[47px] font-bold text-[{ink}]">{e(cfg["title"])}</h1><div class="absolute left-[650px] top-[440px] w-[560px] text-[22px] leading-[32px] text-black">Fall 2026<br>Jordan Lopez, Ph.D.<br>Professor of Neurobiology &amp; Biomedical Engineering</div><div class="absolute left-[60px] top-[625px] w-[1150px] text-[13px] leading-[17px] text-black">{e(title["caption"])}</div>','\n'.join(sp['title_refs'])))
 for sn,s in enumerate(sp['slides'],2):
  f=s['figure'];b=next(sh for sh in prs.slides[sn-1].shapes if sh.name=='Body');pt=b.text_frame.paragraphs[0].runs[0].font.size.pt;px=round(pt*4/3,2);leading=round(px*1.13,2)
  body=''.join(f'<p class="mb-[{round(px*.6,2)}px] text-[{px}px] leading-[{leading}px] text-black">{e(t)}</p>' for t in s['body']);m=by.get(sn);h=221 if m else 440;capy=361 if m else 588
  parts=f'<h2 class="absolute left-[60px] top-[31px] w-[1160px] text-[40px] leading-[46px] font-bold text-[{ink}]">{e(s["title"])}</h2><div class="absolute left-[60px] top-[130px] w-[575px]">{body}</div><img src="{e(url(f["path"]))}" alt="{e(f["caption"])}" width="547" height="{h}" class="absolute left-[673px] top-[134px] h-[{h}px] w-[547px] object-contain"><div class="absolute left-[673px] top-[{capy}px] w-[547px] text-[12.67px] leading-[15px] text-black">{e(f["caption"])}</div>'
  if m:parts+=f'<img src="{e(url(m["poster"]))}" alt="{e(m["caption"])}" width="547" height="202" class="absolute left-[673px] top-[400px] h-[202px] w-[547px] object-contain"><div class="absolute left-[673px] top-[614px] w-[547px] text-[12.67px] leading-[15px] text-black">{e(m["caption"])}</div>'
  parts+=f'<div class="absolute left-[60px] top-[677px] w-[1100px] border-t border-slate-300 pt-[6px] text-[12px] leading-[15px] text-black">{e(s["cite"])}</div><span class="absolute right-[60px] top-[680px] text-[12px] text-black">{sn}</span>'
  slides.append(section(parts,'\n\n'.join(s['transcript'])+'\n\nReferences:\n'+'\n'.join(s['refs'])+'\nStatic companion: source poster; use the native course PPTX for embedded recording playback.'))
 items=''.join(f'<p class="mb-[17px] text-[22px] leading-[27px] text-black"><strong class="text-[{ink}]">{e(t["lead"])}</strong> {e(t["text"])}</p>' for t in sp['takeaways']['items']);slides.append(section(f'<h2 class="absolute left-[60px] top-[35px] text-[40px] font-bold text-[{ink}]">Key takeaways</h2><div class="absolute left-[60px] top-[115px] w-[1150px]">{items}</div>','\n'.join(sp['takeaways']['refs'])))
 doc='<!doctype html><html lang="en"><head><meta charset="utf-8"><script src="https://cdn.tailwindcss.com"></script></head><body class="m-0 p-0"><main id="presentation-slides-wrapper" class="m-0 w-[1280px] p-0">'+''.join(slides)+'</main></body></html>';assert not validate_html(doc),validate_html(doc);assert len(slides)==80
 result={'title':cfg['title'],'slides':80,'formats':{},'fonts':['Arial'],'static_companion':True,'design':'Course-established split column layout with original figures and large source-recording posters. Native course PPTX contains offline videos.'}
 with tempfile.TemporaryDirectory(prefix='presenton-') as td:
  (Path(td)/'presentation.html').write_text(doc)
  for fmt in ['pptx','pdf','png']:
   print(ch,'EXPORT',fmt,flush=True);a=request_json('/api/v3/export/html-to-any',{'html':doc,'format':fmt,'title':cfg['title']+' - Static Companion'},timeout=300);print_response_message(a);assert isinstance(a.get('id'),int) and a['id']>0 and a.get('url','').startswith('https://');result['formats'][fmt]=a;(D/'presenton_export.json').write_text(json.dumps(result,indent=2)+'\n');print(ch,'COMPLETE',fmt,flush=True)
  a=request_json('/api/v3/export/html-to-any/create-preview',{'id':result['formats']['pptx']['id']},timeout=60);print_response_message(a);assert a.get('url','').startswith('https://');result['preview']=a;result['expires']='24 hours';(D/'presenton_export.json').write_text(json.dumps(result,indent=2)+'\n');print_font_usage(doc)
