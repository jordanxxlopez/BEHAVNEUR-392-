"""Check final package, recordings, source crops, notes, and rendered text."""
from pathlib import Path
import json,hashlib,zipfile,posixpath,subprocess,concurrent.futures as C,re
from lxml import etree as E
from pptx import Presentation
import fitz
h=Path(__file__).resolve().parent;s=json.load(open(h/'lecture.json'));m=json.load(open(h/'media_sources.json'))['sources'];p=h/s['output_filename']
def sha(path):
 z=hashlib.sha256()
 with open(path,'rb') as f:
  for b in iter(lambda:f.read(8*1024*1024),b''):z.update(b)
 return z.hexdigest()
with zipfile.ZipFile(p) as z:
 assert z.testzip() is None
 names=set(z.namelist());native={hashlib.sha256(z.read(n)).hexdigest() for n in names if n.startswith('ppt/media/') and n.endswith('.mp4')};assert native=={x['sha256'] for x in m} and len(native)==40
 for x in m:
  assert sha(h/x['file'])==x['sha256'];assert sha(h/x['downloaded_file'])==x['original_sha256'];rel=z.read(f"ppt/slides/_rels/slide{x['slide']}.xml.rels");assert b'/video' in rel and b'/media' in rel;assert b'<p:timing>' in z.read(f"ppt/slides/slide{x['slide']}.xml")
 for n in names:
  if n.endswith('.rels'):
   owner=posixpath.join(posixpath.dirname(posixpath.dirname(n)),posixpath.basename(n)[:-5]) if n!='_rels/.rels' else ''
   for r in E.fromstring(z.read(n)):
    if r.get('TargetMode')!='External':assert posixpath.normpath(posixpath.join(posixpath.dirname(owner),r.get('Target'))).lstrip('/') in names
prs=Presentation(p);assert len(prs.slides)==86
fonts={r.font.name for sl in prs.slides for sh in sl.shapes if sh.has_text_frame for para in sh.text_frame.paragraphs for r in para.runs if r.text.strip()};assert fonts=={'Arial'}
norm=lambda t:re.sub(r'\s+|\*\*','',t)
for sd,sl in zip(s['slides'],list(prs.slides)[1:-1]):
 assert sd['body']==sd['transcript'] and len(sd['body'])==3
 notes=norm(sl.notes_slide.notes_text_frame.text)
 for t in sd['body']:assert norm(t) in notes
 assert not re.search(r'\baesth(?:etic|ic)\b|academic recording|context video',sl.notes_slide.notes_text_frame.text,re.I)
 assert (h/sd['figure']['path']).exists()
for row in json.load(open(h/'article_crops.json')):
 assert sha(h/row['original_image'])==row['original_sha256'];assert sha(h/row['output'])==row['output_sha256']
for row in json.load(open(h/'textbook_crops.json')).values():assert sha(h/row['path'])==row['sha256']
def probe(f):
 a=json.loads(subprocess.check_output(['ffprobe','-v','error','-show_streams','-show_format','-of','json',str(f)],stderr=subprocess.DEVNULL));v=next(x for x in a['streams'] if x['codec_type']=='video');return dict(width=v['width'],height=v['height'],rate=v['r_frame_rate'],average_rate_header=v['avg_frame_rate'],duration=float(a['format']['duration']),audio=[{k:x.get(k) for k in ['codec_name','sample_rate','channels']} for x in a['streams'] if x['codec_type']=='audio'])
def compare(x):
 a=probe(h/x['downloaded_file']);b=probe(h/x['file']);ac=[{k:v for k,v in z.items() if k!='codec_name'} for z in a['audio']];bc=[{k:v for k,v in z.items() if k!='codec_name'} for z in b['audio']];ok=a['rate']==b['rate'] and abs(a['duration']-b['duration'])<=.12 and ac==bc and abs(a['width']-b['width'])<=1 and abs(a['height']-b['height'])<=1;assert ok,x['file'];return dict(file=x['file'],original=a,playable=b,pass_=ok)
with C.ThreadPoolExecutor(4) as ex:probes=list(ex.map(compare,m))
(h/'media_parameter_check.json').write_text(json.dumps(probes,indent=2)+'\n')
a=json.load(open(h/'paragraph_line_check.json'));pdf=h/'The_Rhythmic_Brain_Sleep_and_Biological_Rhythms_Static.pdf';assert a['passing']==252 and not a['failures'];assert a['pptx_sha256']==sha(p) and a['spec_sha256']==sha(h/'lecture.json') and a['pdf_sha256']==sha(pdf)
d=fitz.open(pdf);assert len(d)==86 and d.embfile_count()==0
assert all(d.xref_get_key(i,'Type')[1]!='/EmbeddedFile' for i in range(1,d.xref_length()));assert all(d.xref_get_key(pg.xref,'Annots')[0]=='null' for pg in d);d.close()
videos=json.load(open(h/'media_validation.json'));assert videos['all_full_decodes_pass'] and len(videos['recordings'])==40
roles=[x['instructional_role'] for x in s['slides']];quarters=[sum(q*21<=x['content_index']-1<(q+1)*21 for x in m) for q in range(4)];assert quarters==[10]*4
indices=sorted(x['content_index'] for x in m);gaps=[b-a-1 for a,b in zip([0]+indices,indices+[85])];assert max(gaps)<=3;assert sum(x>=75 for x in indices)>=5
assert [sum(x['purpose']==p for x in m) for p in ['academic','aesthetic']]==[20,20]
frame=next(x['frame_timing_verification'] for x in m if x['source_article']=='PMC4305081');assert frame['original_actual_decoded_frames']==frame['playable_actual_decoded_frames']==4499 and frame['maximum_frame_timestamp_difference_seconds']==0 and frame['audio_packet_hashes_identical']
audit=json.load(open(h/'cross_course_source_audit.json'));assert audit['selected_source_count']==31 and audit['selected_overlaps']==0 and audit['additional_latest_course_check']['overlaps']==[]
v=dict(slides=86,content_slides=84,foundation_slides=roles.count('foundation'),research_slides=roles.count('research'),article_figure_slides=sum(x['figure']['kind']=='article' for x in s['slides']),embedded_distinct_mp4s=40,academic_recordings=20,visual_recordings=20,recordings_per_content_quarter=quarters,maximum_content_gap=max(gaps),recordings_in_final_ten_content_slides=sum(x>=75 for x in indices),paragraphs=252,rendered_paragraph_audit='252/252 pass 5–7 actual rendered lines; hashes identify final files.',all_source_and_embedded_video_hashes_match=True,all_package_relationships_resolve=True,native_video_timing_entries=True,full_video_decode_pass=True,source_dimensions_frame_rates_audio_presence_channels_and_sample_rates_preserved=True,one_pixel_padding_for_odd_source_dimensions=sum(x['width']!=x['original_width'] or x['height']!=x['original_height'] for x in m),h264_transcodes=sum(x['file']!=x['downloaded_file'] for x in m),unchanged_original_mp4s=sum(x['file']==x['downloaded_file'] for x in m),audio_compatibility_conversions='No source audio streams were re-encoded. The converted Haynes MP4 copies all 6461 original AAC packets unchanged.',minimum_conversion_first_60_frames_ssim=min(x['quality_ssim'] for x in videos['recordings'] if x['quality_ssim'] is not None),original_timestamp_notes=[x['file'] for x in videos['recordings'] if x.get('source_timestamp_mux_warning')],timestamp_note='Three original publisher clips have a final-frame timestamp reversal reported by FFmpeg null muxing. Every source frame decodes, and decoded counts equal source counts. Original bytes and timing retained.',static_pdf_pages=86,static_pdf_contains_no_video_attachments=True,fonts=sorted(fonts),selected_primary_papers=31,selected_paper_overlaps_with_checked_courses=0,desktop_powerpoint_playback_tested=False,pptx=dict(filename=p.name,bytes=p.stat().st_size,sha256=sha(p)),static_pdf=dict(filename=pdf.name,bytes=pdf.stat().st_size,sha256=sha(pdf)))
(h/'VALIDATION.json').write_text(json.dumps(v,indent=2)+'\n');print(json.dumps(v,indent=2))
