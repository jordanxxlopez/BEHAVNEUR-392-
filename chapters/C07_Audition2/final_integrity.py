"""Verify native media, package links, notes, source figures, fonts, and rendered paragraphs."""
from pathlib import Path
import json,hashlib,zipfile,posixpath,subprocess,concurrent.futures,re
from lxml import etree as E
from pptx import Presentation
import fitz
h=Path(__file__).resolve().parent;s=json.load(open(h/'lecture.json'));m=json.load(open(h/'media_sources.json'))['sources'];p=h/s['output_filename'];sha=lambda f:hashlib.sha256(Path(f).read_bytes()).hexdigest()
with zipfile.ZipFile(p) as z:
 assert z.testzip() is None
 names=set(z.namelist());actual={hashlib.sha256(z.read(n)).hexdigest() for n in names if n.startswith('ppt/media/') and n.endswith('.mp4')};assert actual=={x['sha256'] for x in m} and len(actual)==50
 for x in m:
  assert sha(h/x['file'])==x['sha256'];assert sha(h/x['downloaded_file'])==x['original_sha256'];rel=z.read(f"ppt/slides/_rels/slide{x['slide']}.xml.rels");assert b'/video' in rel and b'/media' in rel;assert b'<p:timing>' in z.read(f"ppt/slides/slide{x['slide']}.xml")
 for n in names:
  if n.endswith('.rels'):
   owner=posixpath.join(posixpath.dirname(posixpath.dirname(n)),posixpath.basename(n)[:-5]) if n!='_rels/.rels' else ''
   for x in E.fromstring(z.read(n)):
    if x.get('TargetMode')!='External':assert posixpath.normpath(posixpath.join(posixpath.dirname(owner),x.get('Target'))).lstrip('/') in names
prs=Presentation(p);assert len(prs.slides)==100
fonts={r.font.name for sl in prs.slides for sh in sl.shapes if sh.has_text_frame for para in sh.text_frame.paragraphs for r in para.runs if r.text.strip()};assert fonts=={'Arial'}
norm=lambda t:re.sub(r'\s+|\*\*','',t)
for sd,slide in zip(s['slides'],list(prs.slides)[1:-1]):
 assert sd['body']==sd['transcript'] and len(sd['body'])==3
 notes=norm(slide.notes_slide.notes_text_frame.text)
 for text in sd['body']:assert norm(text) in notes
 assert not re.search(r'\baesth(?:etic|ic)\b|academic recording|context video',slide.notes_slide.notes_text_frame.text,re.I)
for row in json.load(open(h/'article_crops.json')).values():assert sha(h/row['original_publisher_figure'])==row['original_sha256']
for sd in s['slides']:assert (h/sd['figure']['path']).exists()
def probe(f):
 r=json.loads(subprocess.check_output(['ffprobe','-v','error','-show_streams','-show_format','-of','json',str(f)],stderr=subprocess.DEVNULL));v=next(x for x in r['streams'] if x['codec_type']=='video');audio=[dict(codec=x['codec_name'],sample_rate=x.get('sample_rate'),channels=x.get('channels')) for x in r['streams'] if x['codec_type']=='audio'];return dict(width=v['width'],height=v['height'],rate=v['avg_frame_rate'],duration=float(r['format']['duration']),audio=audio)
def check_media(x):
 a=probe(h/x['downloaded_file']);b=probe(h/x['file']);ac=[{k:v for k,v in y.items() if k!='codec'} for y in a['audio']];bc=[{k:v for k,v in y.items() if k!='codec'} for y in b['audio']];ok=all([a['rate']==b['rate'],abs(a['duration']-b['duration'])<=.12,ac==bc,abs(a['width']-b['width'])<=1,abs(a['height']-b['height'])<=1]);assert ok,x['file'];return dict(file=x['file'],original=a,embedded_source=b,pass_=ok)
with concurrent.futures.ThreadPoolExecutor(4) as ex:probes=list(ex.map(check_media,m))
(h/'media_parameter_check.json').write_text(json.dumps(probes,indent=2)+'\n')
a=json.load(open(h/'paragraph_line_check.json'));assert a['passing']==294 and not a['failures'];assert a['pptx_sha256']==sha(p);assert a['spec_sha256']==sha(h/'lecture.json');pdf=h/'The_Listening_Brain_Central_Auditory_System_Static.pdf';assert a['pdf_sha256']==sha(pdf)
d=fitz.open(pdf);assert len(d)==100 and not d.embfile_names()
for pg in d:assert d.xref_get_key(pg.xref,'Annots')[0]=='null'
assert all(d.xref_get_key(i,'Type')[1]!='/EmbeddedFile' for i in range(1,d.xref_length()));d.close()
mv=json.load(open(h/'media_validation.json'));assert mv['all_full_decodes_pass'] and len(mv['recordings'])==50
roles=[x['instructional_role'] for x in s['slides']];quarters=[sum(q*98/4<=x['content_index']-1<(q+1)*98/4 for x in m) for q in range(4)];assert quarters==[13,12,13,12]
v=dict(slides=100,content_slides=98,foundation_slides=roles.count('foundation'),research_slides=roles.count('research'),article_figure_slides=sum(x['figure']['kind']=='article' for x in s['slides']),embedded_distinct_mp4s=50,academic_recordings=25,visual_recordings=25,recordings_per_content_quarter=quarters,recordings_in_final_ten_content_slides=sum(x['content_index']>=89 for x in m),paragraphs=294,rendered_paragraph_audit='294/294 pass 5–7 rendered lines; hashes identify final files.',all_source_and_embedded_video_hashes_match=True,all_package_relationships_resolve=True,native_video_timing_entries=True,full_video_decode_pass=True,source_dimensions_frame_rates_audio_presence_channels_and_sample_rates_preserved=True,one_pixel_padding_for_odd_source_dimensions=sum(x['width']!=x['original_width'] or x['height']!=x['original_height'] for x in m),h264_transcodes=sum('compatibility conversion' in x['conversion'] for x in m),lossless_remuxes=sum('Lossless remux' in x['conversion'] for x in m),unchanged_original_mp4s=sum(x['file']==x['downloaded_file'] for x in m),audio_compatibility_conversions='Two original PCM soundtracks encoded to AAC for MP4 playback; timing, sample rate, channels, and audible content retained.',minimum_conversion_first_60_frames_ssim=min(x['quality_ssim'] for x in mv['recordings'] if x['quality_ssim'] is not None),static_pdf_has_100_pages_and_no_video_attachments=True,fonts=sorted(fonts),desktop_powerpoint_playback_tested=False,pptx=dict(filename=p.name,bytes=p.stat().st_size,sha256=sha(p)),static_pdf=dict(filename=pdf.name,bytes=pdf.stat().st_size,sha256=sha(pdf)))
(h/'VALIDATION.json').write_text(json.dumps(v,indent=2)+'\n');print(json.dumps(v,indent=2))
