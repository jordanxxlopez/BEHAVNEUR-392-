from pathlib import Path
import json,copy,shutil,sys,re
from pptx import Presentation
HERE=Path(__file__).resolve().parent;ROOT=HERE.parent.parent;NEURO=Path('/workspace/NEUROETHOLOGY')
sys.path.insert(0,str(HERE));from write_foundations import FOUNDATIONS,BOOK
slides=[];origin=[]
NEW=json.loads((HERE/'new_article_crops.json').read_text());newfig={Path(x['path']).stem:dict(kind='article',path=x['path'],caption=x['caption'],source_url=x['source_url']) for x in NEW}
bookreg=json.loads((HERE/'textbook_crops.json').read_text());bookfigs={Path(x['path']).name:x for x in bookreg['figures']}
def add(title,body,fig,refs,cite,role='research',source='new'):
 sd={'title':title,'body':body,'transcript':body[:],'layout':'figure-right','figure':copy.deepcopy(fig),'figure_width':5.3,'cite':cite,'refs':refs,'instructional_role':role};slides.append(sd);origin.append({'slide':len(slides)+1,'source':source,'role':role});return sd

def foundation(row):
 title,name,body=row;f=copy.deepcopy(bookfigs[name]);f.setdefault('caption',f"Carlson and Birkett (2022), Fig. {f['figure_id']}, p. {f['printed_page']}. "+{'textbook_f1_2_mindbody.png':'Historical accounts of the mind–body relationship.','textbook_f1_7_adaptation.png':'Related structures with different biological functions.','textbook_f1_8_primate_tree.png':'Evolutionary relationships among primates.','textbook_f1_10_brain_counts.png':'Cortical neuron numbers across mammals.'}.get(name,'Original textbook figure.'));add(title,body,f,[BOOK],'Carlson & Birkett (2022)','foundation','Uploaded textbook Chapter 1')

def reuse(n,number,new_title=None):
 folder=NEURO/f'lectures/L{n}';file=folder/'lecture.json'
 if not file.exists():file=folder/'lecture_expanded.json'
 s=copy.deepcopy(json.loads(file.read_text())['slides'][number-2]);s['instructional_role']='research';s['layout']='figures-right' if 'figures' in s else 'figure-right';s['figure_width']=6.0 if 'figures' in s else 5.3
 if 'figures' in s and s.get('figure_arrangement')=='side-by-side':s['primary_figure_width']=4.0
 if new_title:s['title']=new_title
 figs=s.get('figures') or [s['figure']]
 for f in figs:
  src=folder/f['path'];dst=HERE/'figures'/f'L{n}_{src.name}';shutil.copy2(src,dst);f['path']='figures/'+dst.name
 slides.append(s);origin.append({'slide':len(slides)+1,'source':f'NEUROETHOLOGY L{n}, slide {number}, verified study and original figure','role':'research'});return s

HEMI='Kliemann D, Adolphs R, Tyszka JM, Fischl B, Yeo BTT, Nair R, Dubois J, Paul LK (2019). Intrinsic functional connectivity of the brain in adults with a single cerebral hemisphere. Cell Reports 29(8):2398–2407.e4. https://doi.org/10.1016/j.celrep.2019.10.067'
for row in FOUNDATIONS[:3]:foundation(row)
add('A single hemisphere retains organized brain networks',[
'Kliemann and colleagues studied adults who had undergone **hemispherectomy**, surgical removal or disconnection of a cerebral hemisphere, during childhood. Structural magnetic resonance imaging, or **MRI**, documented the remaining brain anatomy. The participants had experienced extensive loss of one side of the cerebrum early in development.',
'The study also measured activity while participants rested in the scanner. A **brain network** is a group of regions whose activity is organized together. Regions within the surviving hemisphere remained organized into recognizable networks associated with vision, movement, attention, and other functions.',
'The remaining tissue retained relationships among regions that are also observed in people with two cerebral hemispheres. The researchers examined how those relationships were arranged within and across networks. These measurements connect childhood changes in anatomy with the organization of adult brain activity.'],newfig['hemi_f1'],[HEMI],'Kliemann et al. (2019)')
add('Functional connectivity measures coordinated fluctuations',[
'**Functional MRI**, or **fMRI**, follows changes in signals associated with blood oxygenation. The **blood-oxygenation-level-dependent** signal is commonly abbreviated **BOLD**. Neural activity and local blood supply are related, allowing researchers to measure slow changes associated with regional brain function.',
'**Functional connectivity** describes the relationship between the changing activity signals of different regions. Kliemann and colleagues compared these signals across cortical areas during rest. Regions assigned to the same network tended to fluctuate together.',
'The cortical map separated visual, movement-related, attention-related, and other networks. A map of regions provided locations for comparing the activity signals. The analysis connected a spatial description of the cortex with the way its signals changed over time.'],newfig['hemi_f2'],[HEMI],'Kliemann et al. (2019)')
add('Brain networks couple more strongly after surgery',[
'The study distinguished relationships within a network from relationships between different networks. Participants with childhood hemispherectomy retained strong within-network organization. Their remaining hemisphere contained coordinated systems rather than a uniformly connected mass of tissue.',
'Relationships between different networks were stronger in almost all hemispherectomy participants. This pattern occurred across multiple network types. Regions belonging to different functional systems therefore had more closely related activity fluctuations.',
'The authors proposed the **hypothesis** that increased interactions between networks characterize functional reorganization after childhood hemispherectomy. The surviving tissue combines recognizable functional systems with altered relationships among those systems. The findings describe a pattern of organization through which researchers can investigate developmental plasticity.'],newfig['hemi_f4'],[HEMI],'Kliemann et al. (2019)')
for row in FOUNDATIONS[3:]:foundation(row)
# Fifteen source-supported optogenetics slides; dedicated control/limitations slides are omitted.
for n in [3,4,5,6,7,8,9,10,12,24,26,42,43,44,15]:reuse(40,n)
# Wireless recording, including recent freely flying bat research.
for n in [32,33,39,40,41]:reuse(40,n)
DLC='Mathis A, Mamidanna P, Cury KM, Abe T, Murthy VN, Mathis MW, Bethge M (2018). DeepLabCut: Markerless pose estimation of user-defined body parts with deep learning. Nature Neuroscience 21:1281–1289. https://doi.org/10.1038/s41593-018-0209-y'
PREPRINT='Original figure crops from the authors’ manuscript, Markerless tracking of user-defined features with deep learning, arXiv:1804.03142 (2018). https://arxiv.org/abs/1804.03142 . Figure numbers explicitly refer to that version.'
dlc=[('Computer vision measures the position of body parts','dlc_f1',[
'**Computer vision** uses computational methods to extract information from images. **Pose estimation** identifies the locations of body parts such as the snout, ears, wrist, or individual digits. A sequence of estimated positions describes how the animal’s configuration changes over time.',
'Mathis and colleagues developed DeepLabCut to locate selected body parts in ordinary video. The user first identifies features of interest in example frames. The trained system then predicts the locations of those features in additional frames.',
'The method was applied to mice following an odor trail, flies moving within a chamber, and mice performing a reaching task. Different features were selected for each preparation. The same analysis framework turned recorded images into measurements tailored to each behavior.']),
('Labeled images define anatomical landmarks','dlc_f2',[
'A **label** specifies the position assigned to a selected body part in an example image. Researchers selected frames containing characteristic postures and marked the locations of the features they wanted to follow. These paired images and labels supplied training information.',
'The image-processing network learned relationships between visual features and body-part locations. The authors trained detectors for features such as the snout and ears. Predictions could then be compared with locations assigned by a human observer.',
'Training used a small collection of labeled images rather than requiring every video frame to be labeled manually. The study reported accurate tracking using about 200 labeled frames in its examples. Once trained, the detector applied the learned relationship to a much longer video sequence.']),
('Transfer learning adapts visual features to a new task','dlc_f1',[
'**Transfer learning** starts with information learned for one computational task and adapts it to another. DeepLabCut used image-processing features from a network previously trained on a large collection of images. The researcher then trained the system to locate body parts relevant to the experiment.',
'An **artificial neural network** is a computational system that transforms inputs through connected processing operations. In this application, the input is an image and the output identifies likely body-part positions. Training changes the computational parameters used to generate those outputs.',
'The authors adapted the approach to both mouse and fly behavior. Reusing previously learned visual features reduced the amount of task-specific labeling required. The resulting measurements followed anatomical landmarks chosen by the researcher.']),
('Tracking resolves head movements during odor following','dlc_f3',[
'During odor-trail following, the mouse moved its snout from side to side across the trail. DeepLabCut tracked the snout together with the ears and tail base. These positions described the movement of the head and the orientation of the body.',
'A **trajectory** is the sequence of positions occupied over time. Combining tracked positions across frames produced snout trajectories during the behavior. The recordings connected the animal’s movement with the location of the odor trail.',
'The study also tracked flies, including body landmarks involved in positioning the head and abdomen. A **proboscis** is the fly’s extendable feeding structure. Tracking its extension provided another behavioral measurement using the same computational approach.']),
('Digit tracking resolves a skilled reaching movement','dlc_f1',[
'A reaching movement combines the motion of the forelimb, wrist, and digits. Mathis and colleagues used a mouse reaching preparation to track selected landmarks across the movement. Their examples included individual digit joints and the wrist.',
'The computer-generated positions allowed the movement to be described at a finer anatomical scale than a single whole-body location. Changes in several landmarks could be measured in the same recorded sequence. The detected positions described the configuration of the hand as the task unfolded.',
'Body-part selection determined the behavioral information extracted from the video. Snout positions supported odor-navigation measurements, while digit positions supported a reaching measurement. Selecting landmarks tied the computational output to the movement that the experiment investigated.'])]
for title,f,b in dlc:add(title,b,newfig[f],[DLC,PREPRINT],'Mathis et al. (2018)')
# Fish-operated vehicle: one core study developed across six slides.
for n in range(2,8):reuse(74,n)
# Sea-lion interfaces, learning, and welfare, including a 2025 follow-up.
for n in [2,3,6,8,17,21,43,40]:reuse(94,n)
# Living networks and feedback-dependent behavior.
for n in [2,3,27,28,29,31]:reuse(71,n)
# Preserve two original Cai et al. figures already cropped in the supplied Biocomputing II deck.
p=Presentation(NEURO/'lectures/L72/Neuroethology_Lecture72_FA2026.pptx')
for number in [35,38]:
 slide=p.slides[number-1];body=[];bodyshapes=[x for x in slide.shapes if x.has_text_frame and len(x.text.split())>30]
 for bodysh in bodyshapes:
  for para in bodysh.text_frame.paragraphs:
   if para.text.strip():body.append(''.join(('**'+r.text+'**' if r.font.bold else r.text) for r in para.runs))
 title=next(x.text for x in slide.shapes if x.has_text_frame);pic=max((x for x in slide.shapes if x.shape_type==13),key=lambda x:x.width*x.height);pic_index=list(slide.shapes).index(pic);cap=next(x.text for x in list(slide.shapes)[pic_index+1:] if x.has_text_frame and x.text.startswith('Cai et al. (2023), Fig.')); file=HERE/'figures'/f'Cai2023_source_slide{number}.{pic.image.ext}';file.write_bytes(pic.image.blob)
 notes=slide.notes_slide.notes_text_frame.text;refs=[line for line in notes.splitlines() if 'Cai' in line and ('2023' in line) and ('doi' in line.lower() or 'Nature Electronics' in line)]
 if not refs:
  refs=['Cai H, Ao Z, Tian C, Wu Z, Liu H, Tchieu J, Gu M, Mackie K, Guo F (2023). Brain organoid reservoir computing for artificial intelligence. Nature Electronics 6:1032–1039. https://doi.org/10.1038/s41928-023-01069-w']
 add(title,body,{'path':'figures/'+file.name,'kind':'article','caption':cap,'source_url':'https://doi.org/10.1038/s41928-023-01069-w'},refs,'Cai et al. (2023)',source=f'Original article crop retained from L72 slide {number}')
# Virtual worlds as action-dependent sensory environments, not fabricated gameplay.
for n in [2,3,4,5,38]:reuse(83,n)
assert len(slides)==73,len(slides)
assert sum(s['instructional_role']=='foundation' for s in slides)==18
# Keep the material in an accessible introductory register with direct titles.
rename={
'Ion substitution identifies the channel’s selectivity':'Different ions carry the light-evoked current',
'Excised membranes retain rapid light gating':'Light gates the channel within the cell membrane',
'Anion channels conduct a different class of ions':'Light can open channels that conduct chloride',
'Optical silencing depends on the chloride gradient':'Chloride movement changes neuronal excitability',
'Position, pulse duty and heating constrain stimulation':'Wireless optical stimulation delivers pulsed energy',
'An organoid acts as an adaptive reservoir':'An organoid transforms a changing electrical input',
'Tectal activation can recruit hunting without retinal input':'Activating a visual region can recruit hunting',
'Large ensembles can be recorded during aerial foraging':'Neural activity can be recorded during free flight',
'Sampling and bandwidth define the observable signal':'Recording systems sample changing neural signals'}
for s in slides:s['title']=rename.get(s['title'],s['title'])
# Replace inherited routine caveat endings with source-supported teaching.
overrides=json.loads((HERE/'teaching_overrides.json').read_text())
for num,change in overrides.items():
 sd=slides[int(num)-2]
 if 'body' in change:sd['body']=change['body']
 else:sd['body'][-1]=change['last']
 sd['transcript']=sd['body'][:]
# Define unfamiliar structures at their first use in this introductory course.
slides[21]['body'][0]='**Channelrhodopsin-2**, or ChR2, is a membrane protein from the green alga _Chlamydomonas reinhardtii_. It contains **retinal**, a light-absorbing molecule. An **ion channel** provides a pathway for charged particles to move across the membrane. Nagel and colleagues expressed ChR2 in frog **oocytes**, immature egg cells, and mammalian cells to investigate its action.'
slides[21]['body'][1]=slides[21]['body'][1].replace('holds membrane voltage at a selected value','holds the electrical potential difference across the cell membrane at a selected value')
slides[21]['body'][-1]+=' Karl Deisseroth, Peter Hegemann, and Georg Nagel received the 2026 Nobel Prize in Physiology or Medicine for optogenetics.'
slides[21]['cite']+='; Graham (2026)'
slides[21]['refs'].append('Graham F (2026). Daily briefing: Medicine Nobel prize goes to optogenetics. Nature, 5 October 2026. https://doi.org/10.1038/d41586-026-03185-x')
slides[31]['body'][0]=slides[31]['body'][0].replace('pretectal neurons','neurons in the **pretectum**, a visual-processing region near the midbrain')
slides[38]['body'][0]=slides[38]['body'][0].replace('dorsal hippocampus','dorsal **hippocampus**, a brain region involved in spatial learning and memory,')
for sd in slides:sd['transcript']=sd['body'][:]
# Introduce the cellular terms used by the biological-computing examples.
sd=slides[60]
sd['body'][0]+=' **Induced pluripotent stem cells** are reprogrammed cells that can develop into many specialized cell types.'
sd['body'][1]=sd['body'][1].replace('supporting astrocytes, a type of glial cell','supporting **astrocytes**, a type of non-neuronal **glial cell**').replace('In these dissociated cultures','In these **dissociated cultures**, where the tissue has been separated into cells,')
sd['body'][1]+=' Dendrites receive inputs, while axons convey outgoing neural signals.'
sd=slides[66]
sd['body'][1]=sd['body'][1].replace('through differentiation and self-organization','through **differentiation**, the development of specialized cellular features, and self-organization')
for sd in slides:sd['transcript']=sd['body'][:]
# Connect pulse control in cultured neurons with the recorded whole-animal application.
mont_ref=next(x for x in slides[34]['refs'] if 'Montgomery' in x)
slides[25]['cite']+='; Montgomery et al. (2015)'
slides[25]['refs'].append(mont_ref)
slides[24]['body'][1]=slides[24]['body'][1].replace(' These controls separate the channel’s electrical effect from fluorescence labeling and illumination.','')
slides[27]['body'][0]=slides[27]['body'][0].replace(' it does not impose a universal hyperpolarizing effect independent of the cell’s ionic conditions.',' the direction of the voltage change follows the ionic conditions in that cell.')
slides[33]['body'][0]=slides[33]['body'][0].replace('A wireless implant must provide enough light while limiting unwanted heating.','Electrical energy powers the light emitter and can warm the surrounding tissue.')
slides[38]['body'][0]=slides[38]['body'][0].replace('wireless Neuropixels recordings','wireless **Neuropixels recordings**, using a probe with many densely packed recording contacts,')
slides[38]['body'][1]=slides[38]['body'][1].replace('single-unit activity','**single-unit activity**, signals assigned to individual neurons,')
slides[60]['body'][1]='The **cortex** is the layered outer region of the cerebral hemispheres. In these **dissociated cultures**, cells separated from the source tissue were dispersed onto a recording surface. Fluorescent labeling identified neuronal cell bodies and their processes. **Dendrites** receive inputs, while **axons** convey outgoing neural signals. **Astrocytes** are a type of non-neuronal glial cell that supports neuronal function.'
for sd in slides:sd['transcript']=sd['body'][:]
# Define the behavioral rating before presenting its relationship with task performance.
sd=slides[59]
sd['body'][0]='Ramos and Winship used ratings from familiar trainers to describe recurring behavioral tendencies. **Personality traits** are patterns in how an individual tends to behave across situations. The study related ratings such as Perceptiveness and Dominance to response-time and button-press measures during computerized tasks.'
sd['body'][1]='Higher Perceptiveness ratings were associated with faster target-contact responses in both initial and later sessions. Higher Dominance ratings were associated with faster initial responses and fewer initial button presses. Response latency and total input count described different aspects of the animal’s performance.'
sd['body'][2]='An animal could differ in response speed, in the number of inputs used, or in both measures. The Perceptiveness association persisted across the training period. The study connected behavioral ratings with recorded performance in a defined cognitive task, while measuring the dimensions of that performance separately.'
for sd in slides:sd['transcript']=sd['body'][:]
# Teach the welfare study’s measured feeding outcomes, not only its apparatus.
sd=slides[58]
sd['title']='Feeding records change after cognitive enrichment'
sd['body'][0]='Winship and colleagues compared feeding records from two-year periods before and after EVE was introduced. **Incomplete consumption of the offered diet** meant that the animal left part of its planned ration unfinished. The records supplied a daily behavioral outcome separate from the response time measured in a computer trial.'
sd['body'][1]='Incomplete-consumption days declined from 61 to 18 for REX and from 40 to 15 for SLD after EVE began. The analysis distinguished feeding records from days removed because of clinical illness. These counts described how often each animal left an unfinished ration within the analyzed periods.'
sd['body'][2]='The average percentage of the offered diet consumed also rose, from about 97% to 99% for REX and from about 98% to 99% for SLD. More complete daily intake and fewer unfinished-ration days described related changes in feeding. The study connected the enrichment period with repeated observations of participation and food consumption in managed animals.'
for sd in slides:sd['transcript']=sd['body'][:]
# Keep the awarded discovery factual and distinct from the experimental findings.
idx=next(i for i,s in enumerate(slides) if 'Membrane expression' in s['title']);slides[idx]['transcript'].append('The 2026 Nobel Prize in Physiology or Medicine recognized Karl Deisseroth, Peter Hegemann, and Georg Nagel for optogenetics, according to Nature’s October 5 report. The experiments on this slide concern optical control of neuronal activity.')
slides[idx]['refs'].append('Graham F (2026). Daily briefing: Medicine Nobel prize goes to optogenetics. Nature, 5 October 2026. https://doi.org/10.1038/d41586-026-03185-x')
refs=list(dict.fromkeys(x for s in slides for x in s['refs']))
spec={'lecture':1,'content_slides':73,'theme':'smoked-steel-paper','output_filename':'From_Brain_to_Behavior_Foundations_and_Frontiers.pptx','title_height':2.4,'title_image':dict(bookfigs['textbook_f1_5_cajal.png']),'title_refs':[BOOK],'slides':slides,'takeaways':{'items':[{'lead':'Biological explanation','text':'Generalization identifies recurring relationships in behavior; reduction explains the physiological processes that produce them.'},{'lead':'Plasticity and development','text':'Experience and development change neural organization, including relationships among networks after early brain injury.'},{'lead':'Evolution and ethics','text':'Inherited variation and natural selection shape behavior; human and animal research includes defined protections and institutional oversight.'},{'lead':'Optical circuit control','text':'Light-sensitive proteins translate illumination into membrane currents, allowing targeted changes in neuronal activity and behavior.'},{'lead':'Behavioral measurement','text':'Neural recordings and computer vision connect measured signals with the animal’s movements and environment.'},{'lead':'Closed-loop learning','text':'A control mapping links an action to a changed input; feedback can alter future behavior in whole animals and living neural networks.'}],'cite':'Carlson & Birkett (2022); Kliemann et al. (2019); Boyden et al. (2005); Mathis et al. (2018); Givon et al. (2022); Kagan et al. (2022)','refs':refs}}
(HERE/'lecture.json').write_text(json.dumps(spec,indent=2,ensure_ascii=False)+'\n');(HERE/'slide_sources.json').write_text(json.dumps(origin,indent=2)+'\n');print('Wrote 75-slide specification; 18 foundations and 55 research content slides')
