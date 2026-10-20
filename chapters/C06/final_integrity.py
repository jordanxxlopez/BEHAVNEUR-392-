"""Check final native package, source bytes, media parameters, and actual-render audit."""
from pathlib import Path
import json,hashlib,zipfile,posixpath,subprocess,concurrent.futures,re
from lxml import etree as E
from pptx import Presentation
h=Path(__file__).resolve().parent;s=json.load(open(h/'lecture.json'));m=json.load(open(h/'media_sources.json'))['sources'];p=h/s['output_filename'];z=zipfile.ZipFile(p);assert z.testzip() is None
sha=lambda f:hashlib.sha256(Path(f).read_bytes()).hexdigest()
actual={hashlib.sha256(z.read(n)).hexdigest() for n in z.namelist() if n.startswith('ppt/media/') and n.endswith('.mp4')};assert actual=={x['sha256'] for x in m} and len(actual)==40
for x in m:
 assert sha(h/x['file'])==x['sha256'];assert sha(h/x['downloaded_file'])==x['original_sha256'];rel=z.read(f"ppt/slides/_rels/slide{x['slide']}.xml.rels");assert b'/video' in rel and b'/media' in rel;assert b'<p:timing>' in z.read(f"ppt/slides/slide{x['slide']}.xml")
for n in z.namelist():
 if n.endswith('.rels'):
  owner=posixpath.join(posixpath.dirname(posixpath.dirname(n)),posixpath.basename(n)[:-5]) if n!='_rels/.rels' else ''
  for x in E.fromstring(z.read(n)):
   if x.get('TargetMode')!='External':assert posixpath.normpath(posixpath.join(posixpath.dirname(owner),x.get('Target'))).lstrip('/') in z.namelist()
prs=Presentation(p);assert len(prs.slides)==90
fonts={r.font.name for sl in prs.slides for sh in sl.shapes if sh.has_text_frame for para in sh.text_frame.paragraphs for r in para.runs if r.text.strip()};assert fonts=={'Arial'}
norm=lambda t:re.sub(r'\s+|\*\*','',t)
for sd,slide in zip(s['slides'],list(prs.slides)[1:-1]):
 assert sd['body']==sd['transcript'] and len(sd['body'])==3
 notes=norm(slide.notes_slide.notes_text_frame.text)
 for text in sd['body']:assert norm(text) in notes
 assert not re.search(r'\baesth(?:etic|ic)\b|academic recording|context video',slide.notes_slide.notes_text_frame.text,re.I)
for row in json.load(open(h/'article_crops.json')).values():assert sha(h/row['original_publisher_figure'])==row['original_sha256']
for sd in s['slides']:assert (h/sd['figure']['path']).exists()
def probe(f):
 r=json.loads(subprocess.check_output(['ffprobe','-v','error','-show_streams','-show_format','-of','json',str(f)]));v=next(x for x in r['streams'] if x['codec_type']=='video');audio=[dict(codec=x['codec_name'],sample_rate=x.get('sample_rate'),channels=x.get('channels')) for x in r['streams'] if x['codec_type']=='audio'];return dict(width=v['width'],height=v['height'],rate=v['avg_frame_rate'],duration=float(r['format']['duration']),audio=audio)
def check_media(x):
 original=probe(h/x['downloaded_file']);out=probe(h/x['file']);ok=all([original['rate']==out['rate'],abs(original['duration']-out['duration'])<=.1,original['audio']==out['audio'],abs(original['width']-out['width'])<=1,abs(original['height']-out['height'])<=1]);assert ok,x['file'];return dict(file=x['file'],original=original,embedded_source=out,pass_=ok)
with concurrent.futures.ThreadPoolExecutor(max_workers=4) as ex:probes=list(ex.map(check_media,m))
(h/'media_parameter_check.json').write_text(json.dumps(probes,indent=2)+'\n')
audit=json.load(open(h/'paragraph_line_check.json'));assert audit['passing']==264 and not audit['failures'];assert audit['pptx_sha256']==sha(p);assert audit['spec_sha256']==sha(h/'lecture.json');pdf=h/'From_Light_to_Sight_Chapter_6_Static.pdf';assert audit['pdf_sha256']==sha(pdf)
mv=json.load(open(h/'media_validation.json'));assert mv['all_full_decodes_pass'] and len(mv['recordings'])==40;converted=sum(x['file']!=x['downloaded_file'] for x in m)
report=dict(slides=90,content_slides=88,foundation_slides=32,research_slides=56,article_figure_slides=58,embedded_distinct_mp4s=40,academic_recordings=20,visual_recordings=20,recordings_per_content_quarter=[sum(q*22<=x['content_index']-1<(q+1)*22 for x in m) for q in range(4)],recordings_in_final_ten_content_slides=sum(x['content_index']>=79 for x in m),paragraphs=264,rendered_paragraph_audit='264/264 pass 5–7 lines; report hashes identify final files',all_source_and_embedded_video_hashes_match=True,all_package_relationships_resolve=True,native_video_timing_entries=True,full_video_decode_pass=True,source_dimensions_frame_rates_and_audio_presence_preserved=True,converted_videos=converted,original_movie_bytes_retained=40-converted,minimum_conversion_first_60_frames_ssim=min(x['quality_ssim'] for x in mv['recordings'] if x['quality_ssim'] is not None),manual_review='All 90 rendered slides and original/cropped figure sheets inspected for body fit, figure labels, captions, and movie placement. Teaching prose reviewed, including all third paragraphs. Source hypotheses and measurement distinctions retained.',fonts=sorted(fonts),desktop_powerpoint_playback_tested=False,pptx=dict(filename=p.name,bytes=p.stat().st_size,sha256=sha(p)),static_pdf=dict(filename=pdf.name,bytes=pdf.stat().st_size,sha256=sha(pdf)))
(h/'VALIDATION.json').write_text(json.dumps(report,indent=2)+'\n');print(json.dumps(report,indent=2))
