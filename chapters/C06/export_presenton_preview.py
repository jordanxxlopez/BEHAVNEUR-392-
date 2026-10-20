"""Export a static Presenton companion; the native PPTX is built separately with embedded movies."""
from pathlib import Path
from pptx import Presentation
import json,hashlib,html,tempfile,urllib.request,time
HERE=Path(__file__).resolve().parent

def px(v):return float(v)/914400*96

def cls(shape):return f'absolute left-[{px(shape.left):.2f}px] top-[{px(shape.top):.2f}px] w-[{px(shape.width):.2f}px] h-[{px(shape.height):.2f}px]'

def post(route,data):
 request=urllib.request.Request('https://api.presenton.ai'+route,data=json.dumps(data).encode(),headers={'Content-Type':'application/json'})
 with urllib.request.urlopen(request,timeout=180) as response:result=json.load(response)
 if result.get('message'):print('API MESSAGE:',result['message'],flush=True)
 return result

def main():
 cache=json.loads((HERE/'presenton_assets.json').read_text());s=json.loads((HERE/'lecture.json').read_text());prs=Presentation(HERE/s['output_filename']);sections=[]
 for slide in prs.slides:
  body=[]
  for shape in slide.shapes:
   c=cls(shape)
   if shape.shape_type==13:
    key=hashlib.sha256(shape.image.blob).hexdigest();url=cache[key]['url'];body.append(f'<img class="{c} object-contain" src="{html.escape(url,quote=True)}" alt="{html.escape(shape.name,quote=True)}"/>')
   elif shape.has_text_frame:
    ps=[]
    for para in shape.text_frame.paragraphs:
     spans=[]
     for run in para.runs:
      size=(run.font.size.pt if run.font.size else 18)*4/3
      try:color=str(run.font.color.rgb)
      except Exception:color='000000'
      if color in ['None','']:color='000000'
      weight='font-bold' if run.font.bold else 'font-normal';italic='italic' if run.font.italic else 'not-italic'
      spans.append(f'<span class="text-[{size:.2f}px] text-[#{color}] {weight} {italic}">{html.escape(run.text)}</span>')
     bottom=(para.space_after.pt if para.space_after else 0)*4/3
     align='text-right' if para.alignment is not None and int(para.alignment)==3 else 'text-left'
     ps.append(f'<p class="m-0 mb-[{bottom:.2f}px] leading-[1.08] {align}">'+''.join(spans)+'</p>')
    body.append(f'<div class="{c} overflow-hidden">'+''.join(ps)+'</div>')
   elif shape.shape_type==1:
    try:color=str(shape.fill.fore_color.rgb)
    except Exception:continue
    body.append(f'<div class="{c} bg-[#{color}]"></div>')
  notes=html.escape(slide.notes_slide.notes_text_frame.text,quote=True)
  sections.append(f'<section class="relative h-[720px] w-[1280px] overflow-hidden bg-white font-[Arial]" data-speaker-note="{notes}">'+''.join(body)+'</section>')
 document='<!doctype html><html lang="en"><head><meta charset="utf-8"><script src="https://cdn.tailwindcss.com"></script></head><body class="m-0 p-0"><main id="presentation-slides-wrapper" class="m-0 w-[1280px] p-0">'+''.join(sections)+'</main></body></html>'
 assert document.count('<section ')==90 and 'data:' not in document and 'style=' not in document
 # TemporaryDirectory ensures HTML cleanup on success or failure; do not persist exported binaries.
 with tempfile.TemporaryDirectory(prefix='presenton-c06-') as temp:
  path=Path(temp)/'presentation.html';path.write_text(document)
  result=post('/api/v3/export/html-to-any',dict(html=document,format='pdf',title='From Light to Sight: Static Lecture Preview'))
  assert isinstance(result.get('id'),int) and result['id']>0 and result.get('url','').startswith(('http://','https://'))
  (HERE/'presenton_export.json').write_text(json.dumps(result,indent=2));print('PDF EXPORT',json.dumps(result),flush=True)
  preview=post('/api/v3/export/html-to-any/create-preview',dict(id=result['id']));assert preview.get('url','').startswith(('http://','https://'));(HERE/'presenton_preview.json').write_text(json.dumps(preview,indent=2));print('PREVIEW',json.dumps(preview),flush=True)
  (HERE/'presenton_fonts.json').write_text(json.dumps(dict(fonts=['Arial'],source='System font; no web-font source',applies_to='Presenton static PDF preview',slides=90),indent=2))
if __name__=='__main__':main()
