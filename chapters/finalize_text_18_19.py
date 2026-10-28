"""Synchronize final teaching text, references, credits, and Arial notes."""
from pathlib import Path
import json,sys
from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor
for ch in map(int,sys.argv[1:]):
 D=Path(__file__).resolve().parent/f'C{ch}';spec=json.load(open(D/'lecture.json'));ms=json.load(open(D/'media_sources.json'))['sources'];by={m['slide']:m for m in ms};p=D/spec['output_filename'];prs=Presentation(p)
 for sn,(slide,sd) in enumerate(zip(list(prs.slides)[1:-1],spec['slides']),2):
  for sh in slide.shapes:
   if sh.name=='Body':
    pt=sh.text_frame.paragraphs[0].runs[0].font.size
    assert len(sh.text_frame.paragraphs)==3
    for para,t in zip(sh.text_frame.paragraphs,sd['body']):
     for r in list(para._p):
      if r.tag.endswith(('}r','}br','}fld')):para._p.remove(r)
     r=para.add_run();r.text=t;r.font.name='Arial';r.font.size=pt;r.font.color.rgb=RGBColor(0,0,0)
   elif sh.name=='Title':sh.text_frame.paragraphs[0].runs[0].text=sd['title']
   elif sh.name=='Figure caption':sh.text_frame.paragraphs[0].runs[0].text=sd['figure']['caption']
   elif sh.name=='Citation':sh.text_frame.paragraphs[0].runs[0].text=sd['cite']
   elif sh.name=='Recording credit':sh.text_frame.paragraphs[0].runs[0].text=by[sn]['caption']
  text='\n\n'.join(sd['transcript'])+'\n\nReferences:\n'+'\n'.join(sd['refs'])
  if sn in by:
   m=by[sn];text+='\n\nWhat students should observe: '+m['teaching_note']+'\nMovie identifier: '+m['movie_identifier']+'\nSource: '+m['original_url']+'\nCredit: '+m['credit']+'\nLicense: '+str(m['license'])+'\nCompatibility: '+m['conversion']+'\nMedia SHA256: '+m['sha256']
  slide.notes_slide.notes_text_frame.text=text
 from PIL import Image
 from pptx.util import Inches
 for sn in ([32,50,74] if ch==19 else []):
  sl=prs.slides[sn-1];f=spec['slides'][sn-2]['figure'];old=next(sh for sh in sl.shapes if sh.shape_type==13);old._element.getparent().remove(old._element);im=Image.open(D/f['path']);ar=im.width/im.height;w=5.7;h=4.7;fw=min(w,h*ar);fh=fw/ar;sl.shapes.add_picture(str(D/f['path']),Inches(7.013+(w-fw)/2),Inches(1.4+(h-fh)/2),Inches(fw),Inches(fh))
 old=next(sh for sh in prs.slides[0].shapes if sh.shape_type==13);x,y,w,h=old.left,old.top,old.width,old.height;im=Image.open(D/spec['title_image']['path']);ar=im.width/im.height;fw=min(w,h*ar);fh=fw/ar;old.left=int(x+(w-fw)/2);old.top=int(y+(h-fh)/2);old.width=int(fw);old.height=int(fh)
 for slide in prs.slides:
  for sh in slide.notes_slide.shapes:
   if sh.has_text_frame:
    for para in sh.text_frame.paragraphs:
     for run in para.runs:run.font.name='Arial';run.font.size=Pt(12)
 prs.save(p);print(ch,'final text and notes synchronized',flush=True)
