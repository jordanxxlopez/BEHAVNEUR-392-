#!/usr/bin/env python3
"""Final course-specific QA, in addition to check_lecture.py.
python tools/check_course.py chapters/C01/lecture.json deck.pptx media_sources.json
"""
import sys,json,zipfile,hashlib,math
from pathlib import Path
from pptx import Presentation

def check(spec,deck,manifest):
 errors=[];slides=spec.get('slides',[]);n=len(slides);total=len(Presentation(deck).slides)
 allowed=spec.get('authorized_total_slide_range',[75,85])
 if not allowed[0]<=total<=allowed[1]:errors.append(f'Total slides {total}; require {allowed}.')
 if total!=n+2:errors.append('Spec/deck slide counts differ.')
 roles=[s.get('instructional_role') for s in slides]
 if any(r not in ('foundation','research') for r in roles):errors.append('Each content slide needs instructional_role foundation or research.')
 if n and not .30<=roles.count('foundation')/n<=.40:errors.append('Foundation slide share must be 30–40%.')
 if n and not .60<=roles.count('research')/n<=.70:errors.append('Research slide share must be 60–70%.')
 if any(not s.get('figure') and not s.get('figures') for s in slides):errors.append('Every content slide needs source imagery.')
 if spec.get('version')==2 and not spec.get('paragraph_line_requirement') and any(len(' '.join(x.get('body',[])).split())>120 for x in slides):errors.append('Version 2 body text exceeds 120 words on a slide.')
 if spec.get('version')==2 and any(__import__('re').search(r'\b(?:connected|linked|supported|relationship)\b',b,__import__('re').I) for x in slides for b in x.get('body',[])):errors.append('Vague summary wording in Version 2 teaching text.')
 if any('textbook' in b.lower() for x in slides for b in x.get('body',[])):errors.append('Visible teaching text must not contain textbook metacommentary.')
 sources=manifest.get('sources',[]);
 policy=json.loads((Path(__file__).resolve().parent.parent/'course/policy.json').read_text())
 policy.update(policy.get('chapter_overrides',{}).get(str(spec.get('lecture')),{}))
 video_min=policy['minimum_total_videos'];video_max=policy['maximum_total_videos']
 if not video_min<=len(sources)<=video_max:errors.append(f'Require {video_min}–{video_max} embedded video records.')
 if spec.get('version')==2 and policy.get('version2_total_videos') is not None and len(sources)!=policy['version2_total_videos']:errors.append(f"Version 2 requires exactly {policy['version2_total_videos']} embedded videos.")
 purposes=[x.get('purpose') for x in sources]
 for purpose in ('academic','aesthetic'):
  minimum,maximum=policy[f'{purpose}_videos']
  if not minimum<=purposes.count(purpose)<=maximum:errors.append(f'Require {minimum}–{maximum} {purpose} video records.')
 if any(p not in ('academic','aesthetic') for p in purposes):errors.append('Each video needs purpose academic or aesthetic.')
 with zipfile.ZipFile(deck) as z:
  embedded={hashlib.sha256(z.read(p)).hexdigest() for p in z.namelist() if p.startswith('ppt/media/') and p.lower().endswith(('.mp4','.mov','.m4v','.avi'))}
 for x in sources:
  if not x.get('original_url') or not x.get('credit'):errors.append(f'Video on slide {x.get("slide")} lacks source URL or credit.')
  if x.get('sha256') not in embedded:errors.append(f'Video on slide {x.get("slide")} is not embedded with its recorded hash.')
  if not isinstance(x.get('slide'),int) or not 2<=x['slide']<total:errors.append('Video slide index must refer to a content slide.')
 if len(set(x.get('sha256') for x in sources))!=len(sources):errors.append('Duplicate movies cannot pad the video quota.')
 if len(embedded)!=len(sources):errors.append('Embedded movie count does not match provenance records.')
 if spec.get('lecture',0)>=policy.get('video_distribution_from_chapter',999):
  indices=sorted(x['slide']-2 for x in sources)
  edges=[-1]+indices+[n]
  if max(b-a-1 for a,b in zip(edges,edges[1:]))>policy['maximum_content_gap_between_videos']:errors.append('Videos leave a gap of more than three content slides.')
  quarters=[sum(q*n/4<=i<(q+1)*n/4 for i in indices) for q in range(4)]
  expected=policy.get('recordings_per_content_quarter',[10,10,10,10])
  if quarters!=expected:errors.append(f'Video distribution by quarter is {quarters}; require {expected}.')
  if sum(i>=n-10 for i in indices)<policy['minimum_videos_in_final_ten_content_slides']:errors.append('Final ten content slides require at least five videos.')
  import re
  for sn,sl in enumerate(Presentation(deck).slides,1):
   visible=' '.join(sh.text for sh in sl.shapes if sh.has_text_frame)
   if re.search(r'\baesth(?:etic|ic)\b|academic recording|context video',visible,re.I):errors.append(f'Visible media-purpose label on slide {sn}.')

 return errors
if __name__=='__main__':
 if len(sys.argv)!=4:sys.exit(__doc__)
 sp=Path(sys.argv[1]);spec=json.loads(sp.read_text());sched=json.loads((Path(__file__).resolve().parent.parent/'course/schedule.json').read_text());meta=next(x for x in sched['lectures'] + sched.get('special_topics',[]) if x['n']==spec['lecture'])
 errors=check(spec,sys.argv[2],json.loads(Path(sys.argv[3]).read_text()))
 if spec.get('paragraph_line_requirement'):
  report_path=sp.parent/'paragraph_line_check.json'
  if not report_path.exists():errors.append('Required actual-render paragraph line audit is missing.')
  else:
   audit=json.loads(report_path.read_text())
   if audit.get('paragraphs')!=sum(len(x['body']) for x in spec['slides']) or audit.get('failures') or audit.get('passing')!=audit.get('paragraphs'):errors.append('Every content paragraph must pass the 5–7-line rendered audit.')
   if audit.get('spec_sha256')!=hashlib.sha256(sp.read_bytes()).hexdigest():errors.append('Line audit does not match the final spec.')
   if audit.get('pptx_sha256')!=hashlib.sha256(Path(sys.argv[2]).read_bytes()).hexdigest():errors.append('Line audit does not match the final PPTX.')
 if not meta.get('verified'):errors.append('Exact chapter title has not been verified against the textbook.')
 for e in errors:print('FAIL:',e)
 print(f'{len(errors)} course-specific failures.')
 sys.exit(bool(errors))
